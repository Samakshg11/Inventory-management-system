from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Color palette matching the glassmorphism dark theme
NAVY = RGBColor(15, 23, 42)          # #0f172a
DARK_BG = RGBColor(30, 41, 59)       # #1e293b
ACCENT = RGBColor(99, 102, 241)      # #6366f1 (Indigo)
ACCENT_LIGHT = RGBColor(129, 140, 248) # #818cf8
TEXT = RGBColor(51, 65, 85)          # #334155
TEXT_LIGHT = RGBColor(148, 163, 184) # #94a3b8
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252)   # #f8fafc
SUCCESS = RGBColor(16, 185, 129)     # #10b981
WARNING = RGBColor(245, 158, 11)     # #f59e0b
BORDER = RGBColor(203, 213, 225)     # #cbd5e1

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_background(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, subtitle_text=""):
        # Top header container
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = NAVY

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(14)
            p2.font.color.rgb = TEXT_LIGHT
            p2.space_before = Pt(4)

        # Bottom divider line
        line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER
        line.line.color.rgb = BORDER

    # ==========================================
    # SLIDE 1: Cover Slide (Dark Glassmorphism)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_background(slide1, NAVY)

    # Accent Glow / Decorative Shape
    glow = slide1.shapes.add_shape(1, Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    glow.fill.solid()
    glow.fill.fore_color.rgb = DARK_BG
    glow.line.color.rgb = ACCENT
    glow.line.width = Pt(2)

    tb1 = slide1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "📦 INVENTORY STOCK MANAGEMENT SYSTEM"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_LIGHT
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf1.add_paragraph()
    p1.text = "A Modern Full-Stack Python, MongoDB & Web SPA Solution"
    p1.font.size = Pt(20)
    p1.font.color.rgb = WHITE
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(15)

    p2 = tf1.add_paragraph()
    p2.text = "College Assignment & Enterprise Grade Architecture"
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_LIGHT
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)

    tb_author = slide1.shapes.add_textbox(Inches(2.0), Inches(5.8), Inches(9.333), Inches(1.0))
    tf_a = tb_author.text_frame
    pa0 = tf_a.paragraphs[0]
    pa0.text = "Developed by: Samaksh Garg"
    pa0.font.size = Pt(16)
    pa0.font.bold = True
    pa0.font.color.rgb = SUCCESS
    pa0.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 2: Project Overview & Problem Statement
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_background(slide2, LIGHT_BG)
    add_header(slide2, "1. Project Overview & Problem Statement", "Addressing retail inefficiencies with a digital, real-time inventory engine")

    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    points2 = [
        ("The Challenge:", "Small retail shops and college stores manually track inventory in paper registers or basic spreadsheets, leading to frequent stockouts, tied-up capital from overstocking, and zero audit trails."),
        ("The Solution:", "A robust, full-stack Inventory Management System featuring a high-performance Python/MongoDB backend, interactive CLI interface, and a modern glassmorphism web dashboard."),
        ("Core Objectives:", "Provide real-time stock tracking, atomic stock operations, comprehensive MongoDB-powered reporting, and seamless data portability via CSV export.")
    ]

    for title, desc in points2:
        p = tf2.add_paragraph()
        p.text = f"• {title} "
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_before = Pt(16)

        p_desc = tf2.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = TEXT
        p_desc.space_before = Pt(4)

    # ==========================================
    # SLIDE 3: Architecture & Tech Stack
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_background(slide3, LIGHT_BG)
    add_header(slide3, "2. Architecture & Technology Stack", "Clean separation of concerns across a 3-tier architecture")

    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    points3 = [
        ("3-Tier Architecture:", "Strict separation across Presentation (CLI/Web SPA), Business Service Layer, and MongoDB Database Layer with singleton connection pooling."),
        ("Python (3.10+):", "Modern type-hinted application core delivering clean, maintainable, PEP 8 compliant code with zero circular dependencies."),
        ("MongoDB 6.x & PyMongo 4.7.3:", "Flexible document model with atomic updates, robust indexing, and powerful aggregation pipelines."),
        ("Supporting Ecosystem:", "Flask & Flask-CORS for REST API, tabulate for CLI tables, faker for realistic seed data, and pytest for comprehensive unit testing.")
    ]

    for title, desc in points3:
        p = tf3.add_paragraph()
        p.text = f"• {title} — {desc}"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_before = Pt(18)

    # ==========================================
    # SLIDE 4: Core Features & Capabilities
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_background(slide4, LIGHT_BG)
    add_header(slide4, "3. Core Features & Capabilities", "Comprehensive toolkit for day-to-day store operations")

    tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    features = [
        ("Product Management:", "Add, paginated view, text/regex search, category filter, update, and delete products by SKU."),
        ("Supplier Management:", "Complete supplier lifecycle with referential integrity checks preventing deletion of active suppliers."),
        ("Stock Operations (IN / OUT):", "Atomic stock updates with guard conditions ensuring zero negative stock."),
        ("Transaction Audit Trail:", "Every stock movement is permanently recorded with timestamp, direction, quantity, and notes."),
        ("CSV Export & Seeding:", "Instant timestamped CSV export and idempotent bulk upsert for 10 suppliers and 40 demo products.")
    ]

    for title, desc in features:
        p = tf4.add_paragraph()
        p.text = f"✓ {title} "
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.space_before = Pt(14)

        p_desc = tf4.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = TEXT
        p_desc.space_before = Pt(4)

    # ==========================================
    # SLIDE 5: Advanced MongoDB Concepts
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_background(slide5, LIGHT_BG)
    add_header(slide5, "4. Advanced MongoDB Concepts Demonstrated", "Leveraging enterprise database features for academic credit")

    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    concepts = [
        ("Atomic $inc Updates:", "Stock IN and OUT use atomic `$inc` operators to guarantee thread-safe operations under concurrency."),
        ("$gte Guard Filter:", "Stock OUT query includes `quantity: {'$gte': requested}` ensuring inventory never drops below zero."),
        ("Indexes & Constraints:", "Unique indexes on `sku` and `email`, text index on product `name`, and compound index on transactions."),
        ("Aggregation Framework:", "Advanced pipelines utilizing `$group`, `$lookup` (joins), `$match`, `$sort`, and `$limit` for rich analytics.")
    ]

    for title, desc in concepts:
        p = tf5.add_paragraph()
        p.text = f"• {title}"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_before = Pt(16)

        p_desc = tf5.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = TEXT
        p_desc.space_before = Pt(4)

    # ==========================================
    # SLIDE 6: Analytics & Aggregation Reports
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_background(slide6, LIGHT_BG)
    add_header(slide6, "5. Analytics & Aggregation Reports", "Five powerful reporting pipelines for business intelligence")

    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    reports = [
        ("1. Low Stock Report:", "Identifies items where quantity <= reorder level for proactive replenishment."),
        ("2. Category-wise Stock Value:", "Groups products by category calculating total quantity and total inventory valuation."),
        ("3. Supplier Summary ($lookup):", "Joins suppliers with products to show product counts and valuation per vendor."),
        ("4. Monthly Stock Movement:", "Aggregates transaction logs by month and movement type (IN vs OUT)."),
        ("5. Top 5 Best-Selling Products:", "Calculates highest total Stock OUT quantity using pipeline sorting and limiting.")
    ]

    for title, desc in reports:
        p = tf6.add_paragraph()
        p.text = f"📊 {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.space_before = Pt(12)

        p_desc = tf6.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = TEXT
        p_desc.space_before = Pt(3)

    # ==========================================
    # SLIDE 7: Live UI — Web Dashboard Screenshot
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_background(slide7, LIGHT_BG)
    add_header(slide7, "6. Live UI: Modern Dashboard", "Glassmorphism SPA dashboard displaying live KPI metrics and stock valuation")

    slide7.shapes.add_picture(
        "/Users/samakshgarg/Desktop/inventory management/docs/screenshots/dashboard.png",
        Inches(1.2), Inches(1.8), width=Inches(10.933)
    )

    # ==========================================
    # SLIDE 8: Live UI — Product Management Screenshot
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_background(slide8, LIGHT_BG)
    add_header(slide8, "7. Live UI: Product Management", "Paginated product catalog with category filtering and real-time search")

    slide8.shapes.add_picture(
        "/Users/samakshgarg/Desktop/inventory management/docs/screenshots/products.png",
        Inches(1.2), Inches(1.8), width=Inches(10.933)
    )

    # ==========================================
    # SLIDE 9: Live UI — Stock Operations & Reports Screenshot
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_background(slide9, LIGHT_BG)
    add_header(slide9, "8. Live UI: Stock Operations & Analytics", "Atomic stock IN/OUT logging and advanced aggregation reports")

    slide9.shapes.add_picture(
        "/Users/samakshgarg/Desktop/inventory management/docs/screenshots/reports.png",
        Inches(1.2), Inches(1.8), width=Inches(10.933)
    )

    # ==========================================
    # SLIDE 10: Testing & Quality Assurance
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_background(slide10, LIGHT_BG)
    add_header(slide10, "9. Testing & Quality Assurance", "Robust test suite ensuring absolute reliability")

    tb10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    testing_pts = [
        ("Comprehensive Test Suite:", "37 unit tests covering product CRUD, stock operations, and aggregation reports."),
        ("Isolated Test Database:", "Tests run against a dedicated `inventory_test_db` with automated fixture teardown ensuring production data is untouched."),
        ("Robust Error Handling:", "Graceful handling of duplicate SKUs, invalid ObjectIds, connection failures, and insufficient stock without raw tracebacks."),
        ("Input Validation:", "Strict validation layer verifying prices > 0, quantities >= 0, and valid email formats before database interaction.")
    ]

    for title, desc in testing_pts:
        p = tf10.add_paragraph()
        p.text = f"✓ {title}"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = SUCCESS
        p.space_before = Pt(16)

        p_desc = tf10.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = TEXT
        p_desc.space_before = Pt(4)

    # ==========================================
    # SLIDE 11: Conclusion & Future Roadmap
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_background(slide11, LIGHT_BG)
    add_header(slide11, "10. Conclusion & Future Roadmap", "Project success and scalable expansion path")

    tb11 = slide11.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    roadmap = [
        ("Project Success:", "Successfully delivered a complete, tested, and polished inventory management system combining CLI and modern web UI."),
        ("Future Enhancement 1:", "Role-based authentication (Admin vs. Operator permissions)."),
        ("Future Enhancement 2:", "Barcode scanner integration for lightning-fast stock check-in/out."),
        ("Future Enhancement 3:", "Automated email/SMS alerts when stock drops below reorder levels."),
        ("Future Enhancement 4:", "Cloud migration from local MongoDB to MongoDB Atlas.")
    ]

    for title, desc in roadmap:
        p = tf11.add_paragraph()
        p.text = f"🚀 {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.space_before = Pt(14)

        p_desc = tf11.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = TEXT
        p_desc.space_before = Pt(3)

    prs.save("Inventory_Management_Presentation.pptx")
    print("Presentation successfully updated with screenshots and polished UI (11 slides).")

if __name__ == "__main__":
    create_presentation()
